#!/usr/bin/env python3
"""生成青海海东智算中心三方合作协商备忘录 Word（公文克制版式）。

体例：A4；上37mm、下35mm、左28mm、右26mm；标题黑体；正文仿宋_GB2312；
一级标题黑体；二级标题楷体_GB2312；三线表；签署栏留空。
"""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, Twips

HERE = Path(__file__).resolve().parent
STEM = "青海海东智算中心节能降碳与综合节能大模型合作备忘录_文秘Grok46_20260908"
TARGET = HERE / f"{STEM}.docx"

FANGSONG = "仿宋_GB2312"
HEITI = "黑体"
KAITI = "楷体_GB2312"
SONGTI = "宋体"

SIZE_TITLE = Pt(18)
SIZE_SUB = Pt(14)
SIZE_H1 = Pt(16)
SIZE_H2 = Pt(16)
SIZE_BODY = Pt(16)
SIZE_TABLE = Pt(10.5)
SIZE_SMALL = Pt(12)
LINE_BODY = Pt(29)


def set_font(run, name: str, size: Pt, bold: bool = False) -> None:
    run.font.name = name
    run.font.size = size
    run.font.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), name)


def add_paragraph(
    doc,
    text: str = "",
    *,
    font: str = FANGSONG,
    size: Pt = SIZE_BODY,
    bold: bool = False,
    align=WD_ALIGN_PARAGRAPH.JUSTIFY,
    first_line_indent: bool = True,
    space_before: Pt = Pt(0),
    space_after: Pt = Pt(0),
    line_spacing: Pt | None = LINE_BODY,
):
    para = doc.add_paragraph()
    fmt = para.paragraph_format
    fmt.alignment = align
    fmt.space_before = space_before
    fmt.space_after = space_after
    if line_spacing is not None:
        fmt.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        fmt.line_spacing = line_spacing
    if first_line_indent:
        fmt.first_line_indent = size * 2
    if text:
        set_font(para.add_run(text), font, size, bold)
    else:
        set_font(para.add_run(""), font, size, bold)
    return para


def set_default_style(doc) -> None:
    style = doc.styles["Normal"]
    style.font.name = FANGSONG
    style.font.size = SIZE_BODY
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), FANGSONG)


def configure_page(section) -> None:
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(37)
    section.bottom_margin = Mm(35)
    section.left_margin = Mm(28)
    section.right_margin = Mm(26)


def add_header(section, text: str) -> None:
    para = section.header.paragraphs[0]
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(para.add_run(text), SONGTI, Pt(10.5))
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "000000")
    borders.append(bottom)
    para._p.get_or_add_pPr().append(borders)


def add_page_number_footer(section) -> None:
    para = section.footer.paragraphs[0]
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def run_with(text: str):
        run = para.add_run(text)
        set_font(run, SONGTI, Pt(12))
        return run

    run_with("— ")
    field = para.add_run()
    set_font(field, SONGTI, Pt(12))
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    field._element.append(begin)
    field._element.append(instr)
    field._element.append(end)
    run_with(" —")


def set_cell_border(cell, **edges) -> None:
    """edges: top/left/bottom/right -> (val, sz, color). val=nil 表示无线。"""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    existing = tcPr.find(qn("w:tcBorders"))
    if existing is not None:
        tcPr.remove(existing)
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        spec = edges.get(edge, ("nil", "0", "000000"))
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), spec[0])
        el.set(qn("w:sz"), spec[1])
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), spec[2])
        tcBorders.append(el)
    tcPr.append(tcBorders)


def shade_cell(cell, fill: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def set_table_widths(table, widths_mm: list[float]) -> None:
    table.autofit = False
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else OxmlElement("w:tblPr")
    if tbl.tblPr is None:
        tbl.insert(0, tblPr)
    existing = tblPr.find(qn("w:tblW"))
    if existing is not None:
        tblPr.remove(existing)
    tblW = OxmlElement("w:tblW")
    total = int(sum(Mm(w) for w in widths_mm) / Twips(1))
    tblW.set(qn("w:w"), str(total))
    tblW.set(qn("w:type"), "dxa")
    tblPr.append(tblW)

    grid = tbl.find(qn("w:tblGrid"))
    if grid is not None:
        tbl.remove(grid)
    grid = OxmlElement("w:tblGrid")
    for w in widths_mm:
        gc = OxmlElement("w:gridCol")
        gc.set(qn("w:w"), str(int(Mm(w) / Twips(1))))
        grid.append(gc)
    tbl.insert(1, grid)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Mm(widths_mm[idx])


def three_line_table(doc, rows: list[list[str]], widths_mm: list[float], caption: str) -> None:
    add_paragraph(
        doc,
        caption,
        font=HEITI,
        size=SIZE_SMALL,
        bold=False,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        first_line_indent=False,
        space_before=Pt(12),
        space_after=Pt(6),
        line_spacing=Pt(20),
    )
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    set_table_widths(table, widths_mm)

    last = len(rows) - 1
    thick = ("single", "12", "000000")
    thin = ("single", "6", "000000")
    none = ("nil", "0", "000000")

    for r, row in enumerate(rows):
        for c, text in enumerate(row):
            cell = table.cell(r, c)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = ""
            para = cell.paragraphs[0]
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER if r == 0 else WD_ALIGN_PARAGRAPH.LEFT
            para.paragraph_format.space_before = Pt(2)
            para.paragraph_format.space_after = Pt(2)
            para.paragraph_format.line_spacing = Pt(16)
            para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            set_font(para.add_run(text), HEITI if r == 0 else FANGSONG, SIZE_TABLE, bold=(r == 0))

            top = thick if r == 0 else none
            bottom = thin if r == 0 else (thick if r == last else none)
            set_cell_border(cell, top=top, bottom=bottom, left=none, right=none)
            if r == 0:
                shade_cell(cell, "F2F2F2")

        tr_pr = table.rows[r]._tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit")
        cant.set(qn("w:val"), "true")
        tr_pr.append(cant)
        if r == 0:
            tbl_header = OxmlElement("w:tblHeader")
            tbl_header.set(qn("w:val"), "true")
            tr_pr.append(tbl_header)


def add_h1(doc, text: str) -> None:
    add_paragraph(
        doc,
        text,
        font=HEITI,
        size=SIZE_H1,
        align=WD_ALIGN_PARAGRAPH.LEFT,
        first_line_indent=False,
        space_before=Pt(16),
        space_after=Pt(6),
    )


def add_h2(doc, text: str) -> None:
    add_paragraph(
        doc,
        text,
        font=KAITI,
        size=SIZE_H2,
        align=WD_ALIGN_PARAGRAPH.LEFT,
        first_line_indent=False,
        space_before=Pt(10),
        space_after=Pt(4),
    )


def add_body(doc, text: str) -> None:
    add_paragraph(doc, text)


def add_sign_line(doc, label: str) -> None:
    add_paragraph(
        doc,
        f"{label}____________________",
        font=FANGSONG,
        size=SIZE_BODY,
        align=WD_ALIGN_PARAGRAPH.LEFT,
        first_line_indent=False,
        space_before=Pt(4),
        space_after=Pt(2),
    )


def build() -> Path:
    doc = Document()
    configure_page(doc.sections[0])
    set_default_style(doc)
    add_header(doc.sections[0], "青海海东智算中心节能降碳与综合节能大模型合作备忘录（讨论稿）")
    add_page_number_footer(doc.sections[0])
    doc.core_properties.title = "青海海东智算中心节能降碳与综合节能大模型合作备忘录"
    doc.core_properties.author = "文秘 · Grok 4.6"
    doc.core_properties.subject = "三方合作协商备忘录（原则安排）讨论稿"
    doc.core_properties.category = "协商备忘录"

    add_paragraph(
        doc,
        "青海海东智算中心节能降碳与综合节能大模型",
        font=HEITI,
        size=SIZE_TITLE,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        first_line_indent=False,
        space_after=Pt(2),
        line_spacing=Pt(28),
    )
    add_paragraph(
        doc,
        "合作备忘录",
        font=HEITI,
        size=SIZE_TITLE,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        first_line_indent=False,
        space_after=Pt(8),
        line_spacing=Pt(28),
    )
    add_paragraph(
        doc,
        "（三方合作协商备忘录 · 原则安排 · 讨论稿）",
        font=KAITI,
        size=SIZE_SUB,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        first_line_indent=False,
        space_after=Pt(10),
        line_spacing=Pt(22),
    )
    for party in (
        "南方电网综合能源股份有限公司",
        "广州恒运",
        "青海超算芯科技有限公司",
    ):
        add_paragraph(
            doc,
            party,
            font=FANGSONG,
            size=SIZE_BODY,
            align=WD_ALIGN_PARAGRAPH.CENTER,
            first_line_indent=False,
            line_spacing=Pt(24),
        )
    add_paragraph(
        doc,
        "二〇二六年九月八日",
        font=FANGSONG,
        size=SIZE_SMALL,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        first_line_indent=False,
        space_before=Pt(6),
        space_after=Pt(16),
        line_spacing=Pt(22),
    )

    add_body(
        doc,
        "本备忘录用于锁定三方共同开发综合节能大模型、开展节能降碳场景合作的原则安排，并把南网能源数据安全和生产数据接入要求写入合作界面。",
    )
    add_body(
        doc,
        "本备忘录确认：大模型和Token是旁路合作，不进入五年闭口保底，也不用于补金租缺口。分成比例、数据权、模型权属和训练经费另签专项协议。本备忘录不构成投资承诺，也不改变南网能源直投五类设备、按闭口服务费回收的既有安排。",
    )

    add_h1(doc, "一、目的和效力")
    add_body(
        doc,
        "各方同意，在青海海东工业互联共享算力池项目（项目主体为青海超算芯科技有限公司，或以其为载体新设的项目公司，以下称SPV）上，共建算力运营团队和算力大模型团队，围绕节能降耗重点项目训练综合节能大模型，并探索将低谷或自留算力做成节能场景产品。运营团队负责上架和回款，大模型团队负责模型与场景。两层工作均不替代闭口保底。",
    )
    add_body(
        doc,
        "本备忘录是协商文本，用于统一口径和推进专项协议，不构成已生效的投资、金租或股权安排。对本备忘录未尽事项，以随后签署的数据合作协议、模型开发协议和Token分成协议为准。本备忘录与五类设备投资、金融租赁、股权转让文件冲突时，投资、金租和股权文件优先；涉及南网能源数据、生产监控和生成式人工智能训练的，以南网能源现行数据安全和生产数据接入制度优先。",
    )

    add_h1(doc, "二、合作各方")
    add_body(
        doc,
        "甲方为南方电网综合能源股份有限公司（以下称南网能源）。南网能源直投并持有储能、机柜、空调、动环和电源五类设备，收取闭口服务费，并按合同参与租赁超额分成；不登记为SPV控股股东。",
    )
    add_body(
        doc,
        "乙方为广州恒运（全称及统一社会信用代码待书面确认，以下称恒运）。恒运为SPV拟登记受让80%股权的一方，主导SPV经营，参与建设和运营。本备忘录项下乙方仅为恒运一家。",
    )
    add_body(
        doc,
        "丙方为青海超算芯科技有限公司（以下称超算芯），提供项目主体和属地条件。",
    )
    add_body(
        doc,
        "原股东青海恒芯能源有限公司拟留股不超过20%；陕西大合九嵩信息科技有限公司拟退出。上述两家原股东不作为本备忘录签署方。落地股东（含恒芯留股等）可不超过20%，具体以随后股权文件为准。",
    )

    add_h1(doc, "三、合作内容")
    add_h2(doc, "（一）节能降碳场景")
    add_body(
        doc,
        "合作聚焦综合能源和算力设施的节能降碳。场景由南网能源提出或确认。本备忘录不承诺节能量、碳减排量、Token销量和单价。",
    )
    add_h2(doc, "（二）综合节能大模型")
    add_body(
        doc,
        "综合节能大模型是技术载体，不替代五类设备投资回报。训练、评测须单独立项。未分级、未授权、未脱敏的数据不得直接用于训练。未通过南网能源数据安全技术管控审查前，不得对接业务系统或对外服务。",
    )
    add_h2(doc, "（三）算力使用")
    add_body(
        doc,
        "可以使用低谷或自留算力，但须服从客户租约和金租资产优先。Token对外销售走旁路账户，不进入监管账户保底瀑布，不得用于弥补金租或闭口缺口。",
    )

    add_h1(doc, "四、与闭口服务和金租的边界")
    add_body(
        doc,
        "项目回收依靠五年闭口可用性服务费，并与上架率脱钩。监管账户支付顺序为：金租、闭口服务费、租赁超额分成，剩余归SPV股东。",
    )
    add_body(
        doc,
        "南网能源不为SPV金租债务提供担保、增信、差额补足或回购。节能收益、Token收入和模型成果，不得解释为金租或闭口的备用还款来源。",
    )

    add_h1(doc, "五、数据与信息安全")
    add_body(
        doc,
        "各方适用《中华人民共和国网络安全法》《中华人民共和国数据安全法》《中华人民共和国个人信息保护法》等现行法律，并执行南网能源《数据安全管理细则（2026版）》（Q/CSG-E2153001-2026）和《生产数据接入及质量管理业务指导书》（Q/CSG-E2154002-2026）。",
    )
    add_body(
        doc,
        "启成平台不是默认训练库，不得默认开放接入。确需使用启成数据的，须分类分级、最小必要、书面授权。增量监控优先采用自建系统。发生安全事件的，须立即通知南网能源数字化管理部门及本备忘录各方。",
    )

    add_h1(doc, "六、成果归属和收益")
    add_body(
        doc,
        "本备忘录不预定著作权、专利、数据权、Token权和分成比例，上述事项由专项协议单列。Token收入和模型收入不进入保底支付顺序。股权分红与本合作项下分成不是同一笔款项，不得混同结算。",
    )

    add_h1(doc, "七、下一步和待书面确认事项")
    add_body(
        doc,
        "各方签署后三十日内，提交数据合作协议、模型开发协议和Token分成协议提纲。未完成分类分级和上线审查的，不得正式训练，不得对外宣传已具备产品。",
    )
    add_body(
        doc,
        "待书面确认事项列于表4。表列事项均保持待书面确认状态，本稿不编造具体数字、名单或比例。",
    )

    add_h1(doc, "八、有效期、签署和其他")
    add_body(
        doc,
        "本备忘录自三方盖章之日起生效，有效期十二个月，或至专项协议生效之日止，以较早者为准。任何一方可提前十五日书面通知终止磋商；保密和数据安全义务在终止后继续有效三年。本备忘录一式三份，三方各执一份。",
    )

    doc.add_page_break()
    add_h1(doc, "九、签署栏")
    add_body(doc, "以下签署栏留空，待正式签署时填写。未盖章前，本稿仅为协商讨论稿。")

    add_paragraph(
        doc,
        "甲方",
        font=HEITI,
        size=SIZE_H1,
        align=WD_ALIGN_PARAGRAPH.LEFT,
        first_line_indent=False,
        space_before=Pt(12),
    )
    add_paragraph(
        doc,
        "南方电网综合能源股份有限公司（盖章）",
        first_line_indent=False,
    )
    add_sign_line(doc, "授权代表：")
    add_sign_line(doc, "职务：")
    add_sign_line(doc, "日期：")

    add_paragraph(
        doc,
        "乙方",
        font=HEITI,
        size=SIZE_H1,
        align=WD_ALIGN_PARAGRAPH.LEFT,
        first_line_indent=False,
        space_before=Pt(16),
    )
    add_paragraph(
        doc,
        "广州恒运（全称待书面确认）（盖章）",
        first_line_indent=False,
    )
    add_sign_line(doc, "授权代表：")
    add_sign_line(doc, "职务：")
    add_sign_line(doc, "日期：")

    add_paragraph(
        doc,
        "丙方",
        font=HEITI,
        size=SIZE_H1,
        align=WD_ALIGN_PARAGRAPH.LEFT,
        first_line_indent=False,
        space_before=Pt(16),
    )
    add_paragraph(
        doc,
        "青海超算芯科技有限公司（盖章）",
        first_line_indent=False,
    )
    add_sign_line(doc, "授权代表：")
    add_sign_line(doc, "职务：")
    add_sign_line(doc, "日期：")

    add_h1(doc, "附表")

    three_line_table(
        doc,
        [
            ["主体", "身份", "主要职责"],
            [
                "南网能源",
                "五类设备投资方，不控股",
                "提出场景和数据安全要求；审核训练数据分级；对涉及公司数据或接入系统的模型与平台行使上线审查",
            ],
            [
                "广州恒运",
                "拟受让SPV 80%股权",
                "主导经营和客户交付；与超算芯落实算力调度、回款和现场条件；按专项协议分担模型编制和经费",
            ],
            [
                "超算芯",
                "项目SPV或载体",
                "提供主体、机房和属地条件；配合数据不出境、分区存储和访问控制；不把租户或客户数据视为已获授权",
            ],
            [
                "运营团队",
                "三方共建",
                "上架、可用性调度、客户交付和回款；费用可随闭口服务费核定，不另开保底分成",
            ],
            [
                "大模型团队",
                "三方共建",
                "围绕节能降耗训练综合节能大模型；编制经费、数据权、模型权属单列",
            ],
        ],
        [26, 44, 86],
        "表1 各方分工",
    )

    three_line_table(
        doc,
        [
            ["项目", "是否进入闭口/保底", "口径"],
            ["五类设备可用性服务费", "进入", "保底与上架脱钩，优于分红"],
            ["底座运维", "可随闭口另核", "对照清能智算维保口径，青海适用性待核"],
            ["运营团队日常费用", "可随闭口另核", "不上架不另开保底分成"],
            ["大模型编制、训练、评测经费", "不进入", "专项协议列支，不挤占7500万元设备款"],
            ["Token和模型分成", "不进入", "旁路结算条款待谈，不纳入投资回报测算"],
            ["金租租金及缺口", "不适用", "本合作不顶、不补、不担保"],
        ],
        [48, 38, 70],
        "表2 与闭口边界",
    )

    three_line_table(
        doc,
        [
            ["要求", "执行要点", "制度来源"],
            [
                "分类分级",
                "训练、接入前须完成分类分级；未分级数据不得直接训练",
                "南网能源《数据安全管理细则（2026版）》（Q/CSG-E2153001-2026）",
            ],
            [
                "三同步与一票否决",
                "按制度实行三同步；未通过审查一票否决，不得上线",
                "同上",
            ],
            [
                "生成式人工智能单独管理",
                "训练、评测、对外服务单列管理，不并入默认业务系统",
                "Q/CSG-E2153001-2026",
            ],
            [
                "目的限制、最小必要",
                "按授权目的使用，不得超范围训练或对外提供",
                "《中华人民共和国数据安全法》《中华人民共和国个人信息保护法》",
            ],
            [
                "访问控制及人员管理",
                "按岗位授权，控制人员范围",
                "Q/CSG-E2153001-2026",
            ],
            ["加密、备份、销毁", "按制度实施加密、备份和到期销毁", "同上"],
            [
                "境内存储，默认不出境",
                "默认境内存储、不出境；出境须另循法定和公司程序",
                "同上",
            ],
            ["留痕审计", "访问、训练、导出须留痕，可供审计", "同上"],
            [
                "生产数据与启成接入必要性评估",
                "启成平台不是默认训练库；生产数据与启成均须作接入必要性评估，确需使用须书面授权",
                "Q/CSG-E2153001-2026、《生产数据接入及质量管理业务指导书》（Q/CSG-E2154002-2026）",
            ],
            [
                "数据质量闭环",
                "接入数据须可校验、可追溯、可纠正",
                "Q/CSG-E2154002-2026",
            ],
        ],
        [40, 62, 54],
        "表3 安全要求",
    )

    three_line_table(
        doc,
        [
            ["序号", "事项", "本稿状态", "说明"],
            [
                "1",
                "广州恒运全称、统一社会信用代码",
                "待书面确认",
                "签署页暂用“广州恒运”，不编造全称或代码",
            ],
            [
                "2",
                "场景清单、试点范围及验收标准",
                "待书面确认",
                "场景须由南网能源提出或确认，本稿不列具体场景",
            ],
            [
                "3",
                "训练数据清单及分级",
                "待书面确认",
                "未完成分类分级不得正式训练",
            ],
            [
                "4",
                "启成接入或免接入结论",
                "待书面确认",
                "启成不是默认训练库，不得默认开放",
            ],
            [
                "5",
                "模型权属、数据权、Token分成",
                "待书面确认",
                "不预定比例，另签专项协议",
            ],
            [
                "6",
                "大模型编制、训练、评测经费",
                "待书面确认",
                "专项协议列支，不挤占7500万元设备款",
            ],
            [
                "7",
                "金租融资比例、闭口年费、超额分成",
                "待书面确认",
                "不属于本备忘录锁定范围",
            ],
        ],
        [14, 48, 28, 66],
        "表4 待书面确认事项（不编造）",
    )

    add_paragraph(
        doc,
        "说明：表4所列事项均待书面确认，不得以本讨论稿代替确认。本稿法律合规表述建议墨律过目，不替代法务审查。",
        size=SIZE_SMALL,
        first_line_indent=False,
        space_before=Pt(10),
        line_spacing=Pt(22),
    )

    doc.save(TARGET)
    return TARGET


def main() -> int:
    path = build()
    print(f"已生成 {path}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
